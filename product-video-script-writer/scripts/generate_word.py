from __future__ import annotations

import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from workflow import copy_sha256


class WordGateError(RuntimeError):
    pass


def _safe_name(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", value).strip().rstrip(".")
    return cleaned[:80] or "未命名产品"


def _set_font(run: Any, size: float, bold: bool = False) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _shade(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), fill)


def _border(cell: Any) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")


def _cell_text(cell: Any, value: str, *, header: bool = False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _shade(cell, "D9EAF7" if header else "FFFFFF")
    _border(cell)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    _set_font(paragraph.add_run(value), 9.5 if header else 9.8, header)


def _next_output(output_root: Path, product_name: str, theme: str) -> Path:
    output_root.mkdir(parents=True, exist_ok=True)
    stem = f"{_safe_name(product_name)}-{_safe_name(theme)}"
    version = 1
    while True:
        candidate = output_root / f"{stem}-V{version}.docx"
        if not candidate.exists():
            return candidate
        version += 1


def _latest_brief(workspace: Any, product_id: str) -> dict[str, Any]:
    events = workspace.events_for_product("products", product_id)
    if not events:
        raise WordGateError("产品库中找不到该产品")
    brief = events[-1].get("payload", {}).get("brief")
    if not isinstance(brief, dict):
        raise WordGateError("产品事件缺少 brief")
    return brief


def generate_word(workspace: Any, product_id: str, script: dict[str, Any], output_root: str | Path) -> Path:
    full_copy = str(script.get("full_copy") or "").strip()
    if not full_copy or not workspace.can_generate_word(product_id, full_copy):
        raise WordGateError("当前完整文案尚未得到明确确认，拒绝生成 Word")
    shot_plan = script.get("shot_plan")
    required_columns = ("time_range", "visual_action", "line", "subtitle_hint", "shooting_note")
    if not isinstance(shot_plan, list) or not shot_plan:
        raise WordGateError("确认后脚本必须补齐 shot_plan")
    for item in shot_plan:
        if not isinstance(item, dict) or any(not str(item.get(key) or "").strip() for key in required_columns):
            raise WordGateError("shot_plan 每段必须包含时间段、画面动作、台词、字幕提示和拍摄备注")

    brief = _latest_brief(workspace, product_id)
    product_name = str(brief.get("product_name") or product_id)
    theme = str(script.get("theme") or "拍摄稿")
    destination = _next_output(Path(output_root), product_name, theme)

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(title.add_run(f"{product_name} 拍摄执行稿"), 20, True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(subtitle.add_run(theme), 13, True)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(meta.add_run(f"形式 {script.get('format')}    预计时长 约 {script.get('estimated_seconds')} 秒"), 9.5)

    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [2.2, 4.1, 6.0, 3.0, 3.0]
    headers = ["时间段", "画面动作", "台词", "字幕提示", "拍摄备注"]
    for index, label in enumerate(headers):
        table.rows[0].cells[index].width = Cm(widths[index])
        _cell_text(table.rows[0].cells[index], label, header=True)
    for item in shot_plan:
        row = table.add_row()
        values = [item[key] for key in required_columns]
        for index, value in enumerate(values):
            row.cells[index].width = Cm(widths[index])
            _cell_text(row.cells[index], str(value))

    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(10)
    _set_font(heading.add_run("道具与场景清单"), 13, True)
    for label, items in (("道具", script.get("props", [])), ("场景", script.get("scenes", []))):
        paragraph = document.add_paragraph(style="List Bullet")
        _set_font(paragraph.add_run(f"{label}：{'、'.join(str(item) for item in items)}"), 10.5)
    copy_heading = document.add_paragraph()
    copy_heading.paragraph_format.space_before = Pt(8)
    _set_font(copy_heading.add_run("完整台词"), 13, True)
    copy_paragraph = document.add_paragraph()
    copy_paragraph.paragraph_format.line_spacing = 1.35
    _set_font(copy_paragraph.add_run(full_copy), 11)

    document.save(destination)
    workspace.append_event("used", {
        "schema_version": 1,
        "event_id": f"evt-{uuid.uuid4()}",
        "event_type": "word_generated",
        "occurred_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "product_id": product_id,
        "user_quote": None,
        "payload": {"candidate_id": script.get("candidate_id"), "path": str(destination), "copy_sha256": copy_sha256(full_copy)},
    })
    return destination
