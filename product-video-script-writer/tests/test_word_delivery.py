from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

TEST_ROOT = Path(__file__).resolve().parent
SKILL_ROOT = TEST_ROOT.parent
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
sys.path.insert(0, str(TEST_ROOT))

from fixtures import brief, packet

from generate_word import WordGateError, generate_word
from workflow import Workspace


def final_script(copy_text: str) -> dict:
    return {
        "candidate_id": "m1",
        "format": "monologue",
        "theme": "键盘缝里的下午茶",
        "estimated_seconds": 30,
        "full_copy": copy_text,
        "props": ["测试专用虚构产品净桌便携清洁器 X1", "键盘", "饼干碎屑"],
        "scenes": ["普通办公室工位"],
        "shot_plan": [
            {"time_range": "0-3 秒", "visual_action": "桌面碎屑特写，产品自然入镜", "line": "每次在工位吃完饼干，桌面好擦，键盘缝最难办。", "subtitle_hint": "键盘缝最难办", "shooting_note": "手机竖屏，近景"},
            {"time_range": "3-22 秒", "visual_action": "按下一键启动并沿键盘移动", "line": "我现在顺手拿它一按就启动，沿着缝走一遍，碎屑很快集中起来。", "subtitle_hint": "一键启动", "shooting_note": "手部特写"},
            {"time_range": "22-30 秒", "visual_action": "展示清理后的桌面", "line": "别等下班再拍键盘，吃完当场收拾，桌面看着也舒服。", "subtitle_hint": "吃完顺手收拾", "shooting_note": "中景自然收尾"},
        ],
    }


class WordDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.workspace = Workspace(self.root / "workspace")
        self.workspace.initialize()
        self.workspace.register_product(brief(), "用户原话：测试产品资料")
        self.copy = packet()["candidates"][0]["full_copy"]
        self.script = final_script(self.copy)

    def test_word_is_refused_before_current_copy_confirmation(self):
        self.workspace.record_selection("test-cleaner-x1", "m1", self.copy, "用户原话：选第一条")
        with self.assertRaisesRegex(WordGateError, "确认"):
            generate_word(self.workspace, "test-cleaner-x1", self.script, self.root / "output")

    def test_confirmed_word_is_editable_and_contains_required_shooting_columns(self):
        self.workspace.record_selection("test-cleaner-x1", "m1", self.copy, "用户原话：选第一条")
        self.workspace.record_confirmation("test-cleaner-x1", self.copy, "用户原话：这版确认")
        path = generate_word(self.workspace, "test-cleaner-x1", self.script, self.root / "output")
        document = Document(path)
        self.assertTrue(any("键盘缝里的下午茶" in p.text for p in document.paragraphs))
        headers = [cell.text for cell in document.tables[0].rows[0].cells]
        self.assertEqual(headers, ["时间段", "画面动作", "台词", "字幕提示", "拍摄备注"])
        self.assertGreater(sum(len(row.cells[2].text) for row in document.tables[0].rows[1:]), 40)

    def test_word_versions_never_overwrite_existing_file(self):
        self.workspace.record_selection("test-cleaner-x1", "m1", self.copy, "用户原话：选第一条")
        self.workspace.record_confirmation("test-cleaner-x1", self.copy, "用户原话：这版确认")
        first = generate_word(self.workspace, "test-cleaner-x1", self.script, self.root / "output")
        second = generate_word(self.workspace, "test-cleaner-x1", self.script, self.root / "output")
        self.assertNotEqual(first, second)
        self.assertTrue(first.exists())
        self.assertTrue(second.exists())
        self.assertTrue(first.name.endswith("-V1.docx"))
        self.assertTrue(second.name.endswith("-V2.docx"))


if __name__ == "__main__":
    unittest.main()
