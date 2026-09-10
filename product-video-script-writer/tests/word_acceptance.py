from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))

from fixtures import brief, packet
from generate_word import WordGateError, generate_word
from test_word_delivery import final_script
from workflow import Workspace


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True, type=Path)
    args = parser.parse_args()
    workspace = Workspace(args.root / "workspace")
    workspace.register_product(brief(), "用户原话：登记测试产品")
    copy_text = packet()["candidates"][0]["full_copy"]
    script = final_script(copy_text)
    workspace.record_selection("test-cleaner-x1", "m1", copy_text, "用户原话：选择第一条")
    try:
        generate_word(workspace, "test-cleaner-x1", script, args.root / "output")
    except WordGateError as exc:
        print(f"PRE_CONFIRM_BLOCKED={exc}")
    else:
        raise AssertionError("Word generation was not blocked before confirmation")
    workspace.record_confirmation("test-cleaner-x1", copy_text, "用户原话：确认当前全文")
    first = generate_word(workspace, "test-cleaner-x1", script, args.root / "output")
    second = generate_word(workspace, "test-cleaner-x1", script, args.root / "output")
    document = Document(first)
    headers = [cell.text for cell in document.tables[0].rows[0].cells]
    assert headers == ["时间段", "画面动作", "台词", "字幕提示", "拍摄备注"]
    assert first.name.endswith("-V1.docx") and second.name.endswith("-V2.docx")
    print(f"V1={first.resolve()}")
    print(f"V2={second.resolve()}")
    print(f"DOCX_STRUCTURE=PASS PARAGRAPHS={len(document.paragraphs)} TABLES={len(document.tables)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
