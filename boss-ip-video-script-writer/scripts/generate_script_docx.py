#!/usr/bin/env python3
"""Generate a versioned editable Word document for one boss-IP shooting script."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from content_contract import duration_label
from history_manager import read_feedback, validate_feedback_event

DOCX_AVAILABLE = True
try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError as exc:
    DOCX_AVAILABLE = False


INK = "243447"
RED = "B23A48"
GOLD = "B7791F"
PALE_RED = "FCECEE"
PALE_GOLD = "FFF7E6"
PALE_BLUE = "EEF4F8"
GRAY = "667085"
LINE = "D7DEE5"
WHITE = "FFFFFF"


class StandardLibraryDocument:
    """Small valid DOCX fallback used when python-docx is unavailable."""

    def __init__(self, data: dict[str, Any]) -> None:
        self.data = data

    def save(self, path: Path | str) -> None:
        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        meta = self.data.get("meta", {})
        paragraphs = [
            clean(self.data.get("cover_title")),
            clean(self.data.get("title")),
            clean(meta.get("estimated_duration")),
            clean(meta.get("scene")),
            "人物",
        ]
        for item in self.data.get("cast", []):
            if isinstance(item, dict):
                paragraphs.append(
                    "｜".join(
                        clean(item.get(key)) for key in ("role", "visibility", "purpose")
                    )
                )
        paragraphs.append("准备内容")
        for item in self.data.get("preparation", []):
            if isinstance(item, dict):
                paragraphs.append(
                    "｜".join(
                        clean(item.get(key)) for key in ("category", "items", "notes")
                    )
                )
        paragraphs.append("分镜执行表")
        for item in self.data.get("segments", []):
            if isinstance(item, dict):
                paragraphs.append(
                    "｜".join(
                        clean(item.get(key))
                        for key in (
                            "time",
                            "stage",
                            "speaker",
                            "dialogue",
                            "visual_action",
                            "camera",
                            "screen_audio",
                        )
                    )
                )
        paragraphs.append("完整对话")
        for item in self.data.get("full_dialogue", []):
            if isinstance(item, dict):
                paragraphs.append(
                    f"{clean(item.get('speaker'))}：{clean(item.get('line'))}"
                )
        paragraphs.extend(["拍摄提醒", *[clean(x) for x in self.data.get("shooting_notes", [])]])
        body = "".join(
            "<w:p><w:r><w:t xml:space=\"preserve\">"
            + escape(text)
            + "</w:t></w:r></w:p>"
            for text in paragraphs
        )
        document_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body>{body}<w:sectPr/></w:body></w:document>"
        )
        content_types = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>"
        )
        relationships = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/>'
            "</Relationships>"
        )
        with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", content_types)
            archive.writestr("_rels/.rels", relationships)
            archive.writestr("word/document.xml", document_xml)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--feedback-ledger", type=Path)
    parser.add_argument("--output-root", type=Path, default=Path("生成脚本") / "发哥老板IP")
    parser.add_argument("--registry", type=Path)
    parser.add_argument("--authorization", type=Path)
    parser.add_argument("--persona", type=Path)
    parser.add_argument("--workspace-root", type=Path, default=Path.cwd())
    return parser.parse_args()


def read_json(path: Path) -> dict[str, Any]:
    try:
        with path.open("r", encoding="utf-8-sig") as handle:
            data = json.load(handle)
    except FileNotFoundError as exc:
        raise SystemExit(f"Input file not found: {path}") from exc
    except json.JSONDecodeError as exc:
        raise SystemExit(f"Invalid JSON at line {exc.lineno}: {exc.msg}") from exc
    if not isinstance(data, dict):
        raise SystemExit("JSON root must be an object.")
    return data


def clean(value: Any, fallback: str = "无") -> str:
    rendered = " ".join(str(value or "").split())
    return rendered or fallback


def safe_component(value: Any, fallback: str) -> str:
    rendered = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", clean(value, fallback)).rstrip(" .")
    return rendered[:80] or fallback


def compact_topic(value: Any) -> str:
    rendered = safe_component(value, "未命名主题")[:16].rstrip(" .-")
    return rendered or "未命名主题"


def confirmation_content_sha256(data: dict[str, Any]) -> str:
    if data.get("schema_version") == 5:
        candidate = data.get("candidate")
        script = data.get("script")
        candidate = candidate if isinstance(candidate, dict) else {}
        script = script if isinstance(script, dict) else {}
        lines = script.get("lines")
        ordered_lines = []
        if isinstance(lines, list):
            ordered_lines = [
                {
                    "line_id": item.get("line_id"),
                    "speaker": item.get("speaker"),
                    "text": item.get("text"),
                }
                for item in lines
                if isinstance(item, dict)
            ]
        payload = {
            "candidate_id": candidate.get("candidate_id"),
            "revision_id": candidate.get("revision_id"),
            "duration_mode": data.get("duration_mode"),
            "theme": script.get("theme"),
            "dialogue": ordered_lines,
            "key_actions": script.get("key_actions"),
            "necessary_shots": script.get("necessary_shots"),
        }
    else:
        payload = copy.deepcopy(data)
        workflow = payload.get("workflow")
        if isinstance(workflow, dict):
            workflow["stage"] = "refined"
            workflow.pop("copy_confirmation", None)
        meta = payload.get("meta")
        if isinstance(meta, dict):
            # The generator assigns V1/V2 only when packaging the confirmed copy.
            # That filename/version metadata is not a change to the confirmed script.
            meta["version"] = "auto"
    rendered = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(rendered).hexdigest()


def _require_text(value: Any, path: str) -> str:
    rendered = clean(value, "")
    if not rendered:
        raise SystemExit(f"Word gate blocked: {path} is required.")
    return rendered


def _validate_production_package(data: dict[str, Any]) -> dict[str, Any]:
    package = data.get("production_package")
    if not isinstance(package, dict):
        raise SystemExit("Word gate blocked: production_package is required.")
    mode = _require_text(data.get("duration_mode"), "duration_mode")
    try:
        expected_label = duration_label(mode)
    except ValueError as exc:
        raise SystemExit(f"Word gate blocked: {exc}") from exc
    label = _require_text(
        package.get("estimated_duration_label"),
        "production_package.estimated_duration_label",
    )
    if label != expected_label:
        raise SystemExit(
            "Word gate blocked: production_package.estimated_duration_label "
            f"must be {expected_label}."
        )
    _require_text(package.get("scene"), "production_package.scene")

    object_arrays = {
        "cast": ("role", "visibility", "purpose"),
        "props": ("item", "purpose"),
        "storyboard": ("beat", "action", "shot"),
        "subtitles": ("line_id", "text"),
    }
    for key, required_keys in object_arrays.items():
        rows = package.get(key)
        if not isinstance(rows, list) or not rows:
            raise SystemExit(f"Word gate blocked: production_package.{key} is required.")
        for index, row in enumerate(rows):
            if not isinstance(row, dict):
                raise SystemExit(
                    f"Word gate blocked: production_package.{key}[{index}] must be an object."
                )
            for field in required_keys:
                _require_text(
                    row.get(field), f"production_package.{key}[{index}].{field}"
                )

    script = data.get("script")
    line_rows = script.get("lines") if isinstance(script, dict) else None
    valid_line_ids = {
        clean(row.get("line_id"), "")
        for row in line_rows or []
        if isinstance(row, dict) and clean(row.get("line_id"), "")
    }
    for index, row in enumerate(package["storyboard"]):
        line_ids = row.get("line_ids")
        if not isinstance(line_ids, list) or not line_ids:
            raise SystemExit(
                f"Word gate blocked: production_package.storyboard[{index}].line_ids is required."
            )
        if any(
            not clean(line_id, "") or clean(line_id, "") not in valid_line_ids
            for line_id in line_ids
        ):
            raise SystemExit(
                f"Word gate blocked: production_package.storyboard[{index}].line_ids "
                "must reference current dialogue lines."
            )
    for index, row in enumerate(package["subtitles"]):
        if clean(row.get("line_id"), "") not in valid_line_ids:
            raise SystemExit(
                f"Word gate blocked: production_package.subtitles[{index}].line_id "
                "must reference a current dialogue line."
            )

    actions = package.get("actions")
    if not isinstance(actions, list) or not actions:
        raise SystemExit("Word gate blocked: production_package.actions is required.")
    for index, action in enumerate(actions):
        _require_text(action, f"production_package.actions[{index}]")
    return package


def validate_v5_word_gate(
    data: dict[str, Any], feedback_rows: list[dict[str, Any]]
) -> None:
    """Require explicit selection and a later confirmation of current v5 copy."""
    candidate = data.get("candidate")
    script = data.get("script")
    if not isinstance(candidate, dict) or not isinstance(script, dict):
        raise SystemExit("Word gate requires v5 candidate and script objects.")
    batch_id = _require_text(candidate.get("batch_id"), "candidate.batch_id")
    candidate_id = _require_text(candidate.get("candidate_id"), "candidate.candidate_id")
    revision_id = _require_text(candidate.get("revision_id"), "candidate.revision_id")
    _validate_production_package(data)

    validated_rows = [validate_feedback_event(row) for row in feedback_rows]
    selection_index = next(
        (
            index
            for index, row in enumerate(validated_rows)
            if row.get("event_type") == "candidate_selection"
            and row.get("batch_id") == batch_id
            and row.get("candidate_id") == candidate_id
        ),
        None,
    )
    if selection_index is None:
        raise SystemExit("Word gate blocked: candidate was not explicitly selected.")

    later_confirmations = [
        row
        for row in validated_rows[selection_index + 1 :]
        if row.get("event_type") == "copy_confirmation"
        and row.get("batch_id") == batch_id
        and row.get("candidate_id") == candidate_id
    ]
    if not later_confirmations:
        raise SystemExit("Word gate blocked: no confirmation follows the selection.")
    current_confirmations = [
        row for row in later_confirmations if row.get("revision_id") == revision_id
    ]
    if not current_confirmations:
        raise SystemExit("Word gate blocked: confirmation does not match current revision.")
    expected = confirmation_content_sha256(data)
    if not any(row.get("content_sha256") == expected for row in current_confirmations):
        raise SystemExit("Word gate blocked: content changed after confirmation.")


def _generation_date(data: dict[str, Any]) -> str:
    candidate = data.get("candidate")
    batch_id = candidate.get("batch_id") if isinstance(candidate, dict) else ""
    match = re.search(r"batch-(\d{4})(\d{2})(\d{2})", clean(batch_id, ""))
    return "-".join(match.groups()) if match else "未标日期"


def prepare_v5_render_data(data: dict[str, Any]) -> dict[str, Any]:
    """Map a confirmed v5 source into the legacy renderer without mutating it."""
    package = _validate_production_package(data)
    candidate = data.get("candidate", {})
    script = data.get("script", {})
    topic = data.get("topic_decision", {})
    evidence = data.get("research_evidence", {})
    line_rows = [row for row in script.get("lines", []) if isinstance(row, dict)]
    lines_by_id = {clean(row.get("line_id"), ""): row for row in line_rows}
    subtitles_by_id: dict[str, list[str]] = {}
    for row in package["subtitles"]:
        subtitles_by_id.setdefault(clean(row.get("line_id"), ""), []).append(
            clean(row.get("text"))
        )

    segments = []
    for board in package["storyboard"]:
        line_ids = [clean(value, "") for value in board["line_ids"]]
        selected_lines = [lines_by_id[value] for value in line_ids]
        speakers = list(dict.fromkeys(clean(row.get("speaker")) for row in selected_lines))
        dialogue = "\n".join(
            f"{clean(row.get('speaker'))}：{clean(row.get('text'))}"
            for row in selected_lines
        )
        screen_audio = "\n".join(
            text for line_id in line_ids for text in subtitles_by_id.get(line_id, [])
        ) or "无"
        segments.append(
            {
                "time": clean(board.get("beat")),
                "stage": clean(board.get("beat")),
                "speaker": "、".join(speakers),
                "dialogue": dialogue,
                "visual_action": clean(board.get("action")),
                "camera": clean(board.get("shot")),
                "screen_audio": screen_audio,
            }
        )

    source_rows = []
    for source in evidence.get("sources", []) if isinstance(evidence, dict) else []:
        if not isinstance(source, dict):
            continue
        source_rows.append(
            {
                "title": source.get("title"),
                "source_kind": "web",
                "source_date": source.get("source_date"),
                "retrieved_at": source.get("retrieved_at"),
                "metric_note": source.get("supports"),
                "borrowed_mechanism": source.get("limits"),
                "locator": source.get("url"),
            }
        )

    theme = clean(script.get("theme"), "未命名主题")
    first_line = line_rows[0] if line_rows else {}
    first_board = package["storyboard"][0]
    first_subtitle = package["subtitles"][0]
    necessary_shots = script.get("necessary_shots", [])
    shooting_notes = [clean(value) for value in package["actions"]]
    shooting_notes.extend(clean(value) for value in necessary_shots if clean(value, ""))
    return {
        "meta": {
            "account_name": "发哥",
            "platform": "抖音自然流量",
            "generation_date": _generation_date(data),
            "filename_topic": theme,
            "version": clean(data.get("output_version"), "auto"),
            "estimated_duration": package["estimated_duration_label"],
            "format": "直接对话"
            if data.get("duration_mode") == "direct_dialogue_30"
            else "轻剧情",
            "scene": package["scene"],
        },
        "content_signature": {
            "category": topic.get("priority_category"),
            "topic": topic.get("topic_id"),
            "priority": "用户选中并确认",
        },
        "title": theme,
        "cover_title": theme,
        "content_positioning": (
            f"{package['estimated_duration_label']}｜{clean(package.get('scene'))}"
        ),
        "core_viewpoint": topic.get("fage_choice"),
        "primary_hook": {
            "voiceover": first_line.get("text"),
            "visual": first_board.get("shot"),
            "screen_text": first_subtitle.get("text"),
        },
        "cast": copy.deepcopy(package["cast"]),
        "preparation": [
            {
                "category": "道具",
                "items": row.get("item"),
                "notes": row.get("purpose"),
            }
            for row in package["props"]
        ],
        "segments": segments,
        "full_dialogue": [
            {"speaker": row.get("speaker"), "line": row.get("text")}
            for row in line_rows
        ],
        "comment_prompt": topic.get("audience_response_space"),
        "shooting_notes": shooting_notes,
        "research_summary": "仅用于选题与冲突背景核验，成稿保持原创表达。",
        "research_sources": source_rows,
        "candidate": copy.deepcopy(candidate),
    }


def validate_word_confirmation(data: dict[str, Any]) -> None:
    candidate = data.get("candidate")
    workflow = data.get("workflow")
    if not isinstance(candidate, dict) or not isinstance(workflow, dict):
        raise SystemExit("Word gate requires candidate and workflow objects.")
    confirmation = workflow.get("copy_confirmation")
    if not isinstance(confirmation, dict):
        raise SystemExit("Word gate requires workflow.copy_confirmation.")
    candidate_id = clean(candidate.get("candidate_id"), "")
    if candidate.get("selection_status") != "selected":
        raise SystemExit("Word gate blocked: candidate was not explicitly selected.")
    if clean(workflow.get("selected_candidate_id"), "") != candidate_id:
        raise SystemExit("Word gate blocked: selected candidate id does not match.")
    if workflow.get("stage") != "copy_confirmed" or confirmation.get("confirmed") is not True:
        raise SystemExit("Word gate blocked: current full copy has no explicit confirmation.")
    revision_id = clean(workflow.get("revision_id"), "")
    if not revision_id or clean(confirmation.get("confirmed_revision_id"), "") != revision_id:
        raise SystemExit("Word gate blocked: confirmation does not match current revision.")
    for key in ["confirmed_at", "user_quote", "content_sha256"]:
        if not clean(confirmation.get(key), ""):
            raise SystemExit(f"Word gate blocked: copy_confirmation.{key} is required.")
    expected = confirmation_content_sha256(data)
    if clean(confirmation.get("content_sha256"), "").lower() != expected:
        raise SystemExit("Word gate blocked: content changed after confirmation.")


def validate_minimum(data: dict[str, Any]) -> None:
    validate_word_confirmation(data)
    for key in ["meta", "content_signature", "primary_hook"]:
        if not isinstance(data.get(key), dict):
            raise SystemExit(f"Required object missing: {key}")
    for key in ["cast", "preparation", "segments", "full_dialogue", "shooting_notes", "research_sources"]:
        if not isinstance(data.get(key), list):
            raise SystemExit(f"Required array missing: {key}")
    for key in ["title", "cover_title", "content_positioning", "core_viewpoint"]:
        if not clean(data.get(key), ""):
            raise SystemExit(f"Required field missing: {key}")
    if not data["segments"] or not data["full_dialogue"]:
        raise SystemExit("segments and full_dialogue may not be empty")


def choose_paths(data: dict[str, Any], output_root: Path) -> tuple[Path, Path, str]:
    meta = data["meta"]
    date_text = safe_component(meta.get("generation_date"), "未标日期")
    topic_text = compact_topic(meta.get("filename_topic") or data.get("cover_title") or data.get("title"))
    output_dir = output_root / date_text
    output_dir.mkdir(parents=True, exist_ok=True)
    requested = clean(meta.get("version"), "auto")
    match = re.fullmatch(r"[Vv]?(\d+)", requested)
    number = max(int(match.group(1)), 1) if match else 1
    while True:
        version = f"V{number}"
        docx_path = output_dir / f"{date_text}-发哥老板IP-{topic_text}-{version}.docx"
        json_path = output_dir / f"{date_text}-发哥老板IP-{topic_text}-{version}-源数据.json"
        if not docx_path.exists() and not json_path.exists():
            return docx_path, json_path, version
        number += 1


def set_repeat_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def set_border(cell: Any, color: str = LINE, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_margins(cell: Any) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", 80), ("start", 90), ("bottom", 80), ("end", 90)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_run(run: Any, size: float = 9.2, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def fill_cell(
    cell: Any,
    value: Any,
    *,
    fill: str = WHITE,
    bold: bool = False,
    color: str | None = None,
    size: float = 9.0,
    align: Any | None = None,
) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_shading(cell, fill)
    set_border(cell)
    set_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    if align is not None:
        paragraph.alignment = align
    style_run(paragraph.add_run(clean(value)), size=size, bold=bold, color=color)


def add_hyperlink(paragraph: Any, text: str, url: str) -> None:
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2F75B5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9.2)
    for name, size, color in [("Title", 21, INK), ("Heading 1", 14, INK), ("Heading 2", 11, RED)]:
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(footer.add_run("发哥｜老板IP拍摄脚本"), size=8, color=GRAY)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_heading(text, level=1)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(4)


def add_table(document: Document, headers: list[str], rows: list[list[Any]], widths: list[float], font_size: float = 8.8) -> Any:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_header(header)
    prevent_split(header)
    for index, label in enumerate(headers):
        header.cells[index].width = Cm(widths[index])
        fill_cell(header.cells[index], label, fill=INK, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_split(row)
        for index, value in enumerate(values):
            row.cells[index].width = Cm(widths[index])
            fill_cell(row.cells[index], value, fill=WHITE if row_index % 2 == 0 else PALE_BLUE, size=font_size)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_title(document: Document, data: dict[str, Any]) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(title.add_run(clean(data.get("cover_title"))), size=21, bold=True, color=INK)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(subtitle.add_run(clean(data.get("title"))), size=13, bold=True, color=RED)
    subtitle.paragraph_format.space_after = Pt(7)
    meta = data["meta"]
    signature = data["content_signature"]
    rows = [
        ["账号", meta.get("account_name"), "平台", meta.get("platform")],
        ["日期", meta.get("generation_date"), "版本", meta.get("version")],
        ["时长", meta.get("estimated_duration"), "形式", meta.get("format")],
        ["场景", meta.get("scene"), "题材", signature.get("category")],
        ["内容定位", data.get("content_positioning"), "核心观点", data.get("core_viewpoint")],
    ]
    table = document.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [2.0, 8.2, 2.0, 14.0]
    for values in rows:
        row = table.add_row()
        prevent_split(row)
        for index, value in enumerate(values):
            row.cells[index].width = Cm(widths[index])
            is_label = index % 2 == 0
            fill_cell(row.cells[index], value, fill=PALE_RED if is_label else WHITE, bold=is_label, color=INK if is_label else None)


def add_hook(document: Document, data: dict[str, Any]) -> None:
    add_heading(document, "1. 唯一主钩子")
    hook = data["primary_hook"]
    rows = [[hook.get("voiceover"), hook.get("visual"), hook.get("screen_text")]]
    add_table(document, ["前三秒对白", "第一帧画面", "屏幕大字"], rows, [8.5, 10.5, 7.2], 9.5)


def add_cast_and_preparation(document: Document, data: dict[str, Any]) -> None:
    add_heading(document, "2. 人物与拍摄准备")
    cast_rows = [
        [item.get("role"), item.get("visibility"), item.get("purpose")]
        for item in data["cast"]
        if isinstance(item, dict)
    ]
    add_table(document, ["人物", "出镜方式", "本条作用"], cast_rows, [4.0, 5.0, 17.2])
    prep_rows = [
        [item.get("category"), item.get("items"), item.get("notes")]
        for item in data["preparation"]
        if isinstance(item, dict)
    ]
    add_table(document, ["类别", "准备内容", "状态与要求"], prep_rows, [3.5, 10.0, 12.7])


def add_segments(document: Document, data: dict[str, Any]) -> None:
    add_heading(document, "3. 分镜执行表")
    rows = []
    for item in data["segments"]:
        rows.append(
            [
                f"{clean(item.get('time'))}\n{clean(item.get('stage'))}",
                item.get("speaker"),
                item.get("dialogue"),
                item.get("visual_action"),
                item.get("camera"),
                item.get("screen_audio"),
            ]
        )
    add_table(
        document,
        ["时间/功能", "说话人", "对白", "动作与画面", "景别/机位", "字幕/声音"],
        rows,
        [2.8, 2.6, 7.2, 5.4, 4.4, 3.8],
        8.4,
    )


def add_dialogue(document: Document, data: dict[str, Any]) -> None:
    add_heading(document, "4. 完整对白")
    rows = [
        [item.get("speaker"), item.get("line")]
        for item in data["full_dialogue"]
        if isinstance(item, dict)
    ]
    table = add_table(document, ["说话人", "对白"], rows, [4.0, 22.2], 10.3)
    for row in table.rows[1:]:
        set_shading(row.cells[1], PALE_GOLD)
        set_border(row.cells[1], GOLD, "6")
    comment = clean(data.get("comment_prompt"))
    paragraph = document.add_paragraph()
    style_run(paragraph.add_run("评论互动："), bold=True, color=RED)
    style_run(paragraph.add_run(comment), size=9.5)


def add_notes(document: Document, data: dict[str, Any]) -> None:
    add_heading(document, "5. 拍摄与剪辑提醒")
    for note in data["shooting_notes"]:
        paragraph = document.add_paragraph(style="List Bullet")
        style_run(paragraph.add_run(clean(note)), size=9.3)
        paragraph.paragraph_format.space_after = Pt(1)


def add_research(document: Document, data: dict[str, Any]) -> None:
    add_heading(document, "6. 参考来源与原创改写说明")
    summary = document.add_paragraph()
    style_run(summary.add_run(clean(data.get("research_summary"))), size=9.5)
    headers = ["来源", "来源日期", "数据说明", "本条只借鉴", "定位"]
    rows = []
    for source in data["research_sources"]:
        rows.append(
            [
                f"{clean(source.get('title'))}\n{clean(source.get('source_kind'))}",
                f"来源 {clean(source.get('source_date'))}\n读取 {clean(source.get('retrieved_at'))}",
                source.get("metric_note"),
                source.get("borrowed_mechanism"),
                source.get("locator"),
            ]
        )
    table = add_table(document, headers, rows, [6.0, 4.2, 6.2, 6.0, 3.8], 7.9)
    for row, source in zip(table.rows[1:], data["research_sources"]):
        cell = row.cells[4]
        locator = clean(source.get("locator"))
        if source.get("source_kind") == "web" and locator.startswith(("http://", "https://")):
            cell.text = ""
            set_shading(cell, WHITE)
            set_border(cell)
            set_margins(cell)
            paragraph = cell.paragraphs[0]
            add_hyperlink(paragraph, "打开来源", locator)


def build_document(data: dict[str, Any]) -> Document:
    if not DOCX_AVAILABLE:
        return StandardLibraryDocument(data)
    document = Document()
    configure(document)
    add_title(document, data)
    add_hook(document, data)
    add_cast_and_preparation(document, data)
    add_segments(document, data)
    add_dialogue(document, data)
    add_notes(document, data)
    add_research(document, data)
    return document


def main() -> int:
    args = parse_args()
    source = read_json(args.input)
    if source.get("document_kind") == "boss_ip_delivery":
        from word_v6 import generate
        if args.registry is None or args.persona is None or args.feedback_ledger is None:
            raise SystemExit("v6 Word requires --registry, --persona and --feedback-ledger")
        try:
            word, saved = generate(source, root=args.workspace_root, output_root=args.output_root,
                registry_path=args.registry, persona=read_json(args.persona),
                feedback_rows=read_feedback(args.feedback_ledger),
                authority=read_json(args.authorization) if args.authorization else None)
        except (ValueError, OSError, KeyError, TypeError) as exc:
            print(f"INTEGRITY FAIL: {exc}")
            return 2
        print(f"Created Word: {word.resolve()}")
        print(f"Created source JSON: {saved.resolve()}")
        return 0
    is_v5 = source.get("schema_version") == 5
    if is_v5:
        if args.feedback_ledger is None:
            raise SystemExit("Word gate for schema v5 requires --feedback-ledger.")
        validate_v5_word_gate(source, read_feedback(args.feedback_ledger))
        data = prepare_v5_render_data(source)
    else:
        validate_minimum(source)
        data = copy.deepcopy(source)
    docx_path, json_path, version = choose_paths(data, args.output_root)
    data["meta"]["version"] = version
    document = build_document(data)
    document.save(docx_path)
    if is_v5:
        saved_source = copy.deepcopy(source)
        saved_source["output_version"] = version
    else:
        saved_source = data
    json_path.write_text(
        json.dumps(saved_source, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"Created Word: {docx_path.resolve()}")
    print(f"Created source JSON: {json_path.resolve()}")
    print(f"Version: {version}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
