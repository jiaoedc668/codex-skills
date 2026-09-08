"""Approved V3 typography with complete v6 dialogue and production notes."""
from __future__ import annotations
from datetime import datetime
import json
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import delivery_contract as dc
from content_contract import validate_candidate
from dual_format_contract import FORMATS


def build_document(envelope):
    dc.validate_production(envelope)
    c=envelope["candidate_doc"];s=c["script"];pack=envelope["production_package"];identity=envelope["delivery"]
    d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65)
    sec.left_margin=sec.right_margin=Inches(.8)
    sec.header_distance=sec.footer_distance=Inches(.32)
    for name in ("Normal","Title","Heading 1","Heading 2"):
        st=d.styles[name];st.font.name="微软雅黑";st.font.color.rgb=RGBColor.from_string("222A30")
        st.paragraph_format.space_after=Pt(6)
    for border in list(d.styles.element.iter(qn("w:pBdr"))):border.getparent().remove(border)
    d.styles["Normal"].font.size=Pt(11);d.styles["Normal"].paragraph_format.line_spacing=1.15
    d.styles["Title"].font.size=Pt(25);d.styles["Title"].font.bold=True;d.styles["Title"].font.color.rgb=RGBColor(0,0,0)
    d.styles["Title"].paragraph_format.line_spacing=1.15
    d.styles["Heading 1"].font.size=Pt(13);d.styles["Heading 1"].paragraph_format.space_before=Pt(10)
    def para(text,size=11,color="222A30",bold=False,after=6):
        p=d.add_paragraph();p.paragraph_format.space_after=Pt(after)
        r=p.add_run(text);r.font.size=Pt(size);r.font.color.rgb=RGBColor.from_string(color);r.bold=bold
        return p
    def table(widths):
        t=d.add_table(rows=0,cols=len(widths));t.autofit=False
        for col,width in zip(t.columns,widths):col.width=Inches(width)
        return t
    def row(t,values,fill=None,size=11,last=False):
        cells=t.add_row().cells
        t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index,(cell,value) in enumerate(zip(cells,values)):
            tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement("w:tcMar")
            for edge,amount in (("top",65),("bottom",65),("left",140),("right",140)):
                item=OxmlElement("w:"+edge);item.set(qn("w:w"),str(amount));item.set(qn("w:type"),"dxa");margins.append(item)
            tcpr.append(margins)
            if fill:
                shade=OxmlElement("w:shd");shade.set(qn("w:fill"),fill);tcpr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_together=True
            r=p.add_run(value);r.font.size=Pt(size);r.bold=index==0 or last
            if index==0:r.font.color.rgb=RGBColor.from_string("8C7041")
    label=f'{identity["script_id"]}  /  V{identity["version"]}'
    h=sec.header.paragraphs[0];h.text="发哥  /  老板 IP";h.runs[0].font.size=Pt(8)
    f=sec.footer.paragraphs[0];f.text="拍摄脚本    ·    "+label+"                          ";f.runs[0].font.size=Pt(8)
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");f._p.append(field)
    para(label+"    /    完整对白",10,"8C7041",True,10)
    d.add_paragraph(s["theme"],"Title")
    para(FORMATS[c["format"]][0]+"    /    "+FORMATS[c["format"]][1]+"    /    "+pack["scene"],10,"70787E",False,17)
    t=table([.75,6.15])
    for index,line in enumerate(s["lines"]):
        last=index==len(s["lines"])-1
        row(t,[line["speaker"],line["text"]],"F4EFE5" if last else ("F5F6F6" if index%2==0 else None),12,last)
    if s["key_actions"]:
        para("收尾动作",9,"8C7041",True,3)
        para("；".join(a["action"] for a in s["key_actions"]),10,"646C72",False,0)
    d.add_page_break()
    para(label+"    /    拍摄执行",10,"8C7041",True,10)
    d.add_paragraph("现场执行单","Title")
    para(s["theme"],10,"70787E",False,12)
    t=table([.75,6.15])
    row(t,["场景",pack["scene"]],"F5F6F6")
    row(t,["人物","；".join(x["role"]+" · "+x["visibility"] for x in pack["cast"])])
    row(t,["道具","、".join(pack["props"]) or "无需额外道具"],"F5F6F6")
    d.add_heading("分镜安排",1);t=table([1.2,5.7])
    for index,shot in enumerate(pack["storyboard"]):
        ids=shot["line_ids"];span="第"+ids[0][1:]+"至"+ids[-1][1:]+"句"
        row(t,[f"{index+1:02d}\n"+span,shot["action"]+"\n"+shot["shot"]],"F5F6F6" if index%2==0 else None)
    if s["key_actions"]:
        d.add_heading("关键动作",1)
        for action in s["key_actions"]:para("第"+action["after_line_id"][1:]+"句后  "+action["action"])
    d.add_heading("必要镜头",1);para("；".join(s["necessary_shots"]))
    d.add_heading("字幕",1);para(pack["subtitles"])
    paragraphs=list(d.paragraphs)+[p for t in d.tables for r in t.rows for cell in r.cells for p in cell.paragraphs]
    for section in d.sections:paragraphs.extend(section.header.paragraphs);paragraphs.extend(section.footer.paragraphs)
    for p in paragraphs:
        bold=p.style.name in {"Title","Heading 1","Heading 2"} or any(r.bold is True for r in p.runs)
        for run in p.runs:
            run.font.name="微软雅黑";run.bold=bold;run.font.cs_bold=bold
            fonts=run._element.get_or_add_rPr().rFonts
            for key in list(fonts.attrib):del fonts.attrib[key]
            for attr in ("ascii","hAnsi","eastAsia","cs"):fonts.set(qn("w:"+attr),"微软雅黑")
    d.core_properties.title=s["theme"];d.core_properties.author="";d.core_properties.last_modified_by=""
    return d


def generate(envelope,*,root,output_root,registry_path,persona,feedback_rows,authority=None):
    dc.validate_production(envelope)
    candidate=envelope["candidate_doc"]
    authored_rows=dc.authored_feedback_prefix(candidate,feedback_rows)
    failures=validate_candidate(candidate,persona,authored_rows).failures
    if failures:raise ValueError("; ".join(failures))
    if authority is None:dc.ordinary_gate(candidate,feedback_rows)
    else:dc.special_gate(envelope,authority,root)
    registry_path=Path(registry_path)
    rows=dc.load(registry_path) if registry_path.exists() else []
    sid,version=dc.allocate(rows,candidate["candidate"]["candidate_id"],candidate["format"])
    if envelope["delivery"]!={"script_id":sid,"version":version}:raise ValueError("delivery number/version differs from registry allocation")
    if version>1:dc.verify_lineage(envelope,root)
    elif envelope["lineage"] is not None:raise ValueError("first delivery cannot claim revision lineage")
    name=re.sub(r'[<>:"/\\|?*\x00-\x1f]',"-",candidate["script"]["theme"]).rstrip(" .")
    if not name or len(name)>80:raise ValueError("title not suitable for filename")
    output_root=Path(output_root)
    if not output_root.is_dir():raise ValueError("output directory must be established before generation")
    stem=f"{sid}-{name}-V{version}";word=output_root/(stem+".docx");source=output_root/(stem+"-源数据.json")
    if word.exists() or source.exists():raise ValueError("output already exists; will not overwrite")
    registry_before=registry_path.read_bytes() if registry_path.exists() else None
    lock=registry_path.with_suffix(".lock")
    created=[]
    with lock.open("x",encoding="utf-8") as handle:
        try:
            if (registry_path.read_bytes() if registry_path.exists() else None)!=registry_before:raise ValueError("registry changed during generation")
            build_document(envelope).save(word);created.append(word)
            with source.open("x",encoding="utf-8") as f:f.write(json.dumps(envelope,ensure_ascii=False,indent=2)+"\n")
            created.append(source)
            event={"candidate_id":candidate["candidate"]["candidate_id"],"script_id":sid,"version":version,"format":candidate["format"],"batch_id":candidate["candidate"]["batch_id"],"source_sha256":dc.file_hash(source),"word_sha256":dc.file_hash(word),"source_path":str(source.resolve()),"word_path":str(word.resolve()),"authorization_kind":"special_rewrite" if authority else "ordinary_confirmation","authorization_sha256":dc.content_hash(authority) if authority else dc.content_hash(feedback_rows),"recorded_at":datetime.now().astimezone().isoformat()}
            dc.validate_registry(rows+[event])
            temporary=registry_path.with_suffix(".tmp")
            with temporary.open("x",encoding="utf-8") as f:f.write(json.dumps(rows+[event],ensure_ascii=False,indent=2)+"\n")
            temporary.replace(registry_path)
        except Exception:
            for path in created:path.unlink()
            raise
        finally:
            handle.close();lock.unlink()
    return word,source
