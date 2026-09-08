import importlib
import importlib.util
import tempfile
import unittest
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"scripts"))
from test_dual_format_v6 import candidate_v6
from docx import Document
from docx.oxml.ns import qn

def envelope_v6(mode="story"):
    c=candidate_v6(mode)
    speakers=list(dict.fromkeys(l["speaker"] for l in c["script"]["lines"]))
    return {"document_kind":"boss_ip_delivery","schema_version":1,"candidate_doc":c,"delivery":{"script_id":"B01" if mode=="story" else "A01","version":1},"lineage":None,
    "production_package":{"scene":"办公室休息区","cast":[{"role":s,"visibility":"露脸" if s=="发哥" else "画外音"} for s in speakers],"props":[],"storyboard":[{"line_ids":[l["line_id"] for l in c["script"]["lines"]],"action":"发哥听完后给出建议。" if mode=="story" else "发哥面对镜头说话。","shot":"固定中景"}],"subtitles":"完整对白逐句同步，按自然停顿分行。"}}

class WordV6Tests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(importlib.util.find_spec("word_v6"),"v6 Word generator not implemented")
        self.m=importlib.import_module("word_v6")

    def test_build_contains_full_copy_number_and_version(self):
        for mode in ("story","monologue"):
            e=envelope_v6(mode);d=self.m.build_document(e)
            paragraphs=list(d.paragraphs)+[p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]
            text="\n".join(p.text for p in paragraphs)
            for line in e["candidate_doc"]["script"]["lines"]:self.assertIn(line["text"],text)
            self.assertIn(e["delivery"]["script_id"],text);self.assertIn("V1",text)
            self.assertIn("口播" if mode=="monologue" else "剧情",text)
            for p in paragraphs:
                for run in p.runs:
                    fonts=run._element.get_or_add_rPr().rFonts
                    for attr in ("ascii","hAnsi","eastAsia","cs"):self.assertEqual("微软雅黑",fonts.get(qn("w:"+attr)))

    def test_production_rejects_extra_monologue_actor_and_unknown_line(self):
        import delivery_contract as dc
        e=envelope_v6("monologue");dc.validate_production(e)
        e["production_package"]["cast"].append({"role":"朋友","visibility":"画外音"})
        with self.assertRaises(ValueError):dc.validate_production(e)
        e=envelope_v6();e["production_package"]["storyboard"][0]["line_ids"]=["L999"]
        with self.assertRaises(ValueError):dc.validate_production(e)

class WordCliV6Tests(unittest.TestCase):
    def test_real_cli_confirmation_gate_generation_and_no_overwrite(self):
        import json
        import subprocess
        import delivery_contract as dc
        from v5_fixtures import valid_persona
        from history_manager import feedback_context
        scripts=Path(__file__).resolve().parents[1]/"scripts"
        for mode in ("story","monologue"):
            with self.subTest(mode=mode),tempfile.TemporaryDirectory() as tmp:
                root=Path(tmp);e=envelope_v6(mode);c=e["candidate_doc"];i=c["candidate"]
                persona=root/"persona.json";persona.write_text(json.dumps(valid_persona()),encoding="utf-8")
                ledger=root/"feedback.jsonl";ledger.write_text("",encoding="utf-8")
                source=root/"input.json";source.write_text(json.dumps(e),encoding="utf-8")
                registry=root/"registry.json"
                cmd=[sys.executable,str(scripts/"generate_script_docx.py"),"--input",str(source),"--registry",str(registry),"--persona",str(persona),"--feedback-ledger",str(ledger),"--output-root",str(root)]
                denied=subprocess.run(cmd,capture_output=True,text=True,encoding="utf-8")
                self.assertEqual(2,denied.returncode,denied.stdout+denied.stderr)
                self.assertEqual([],list(root.glob("*.docx")));self.assertFalse(registry.exists())
                selection={"event_type":"candidate_selection","feedback_id":"sel","batch_id":i["batch_id"],"candidate_id":i["candidate_id"],"format":mode,"user_quote":"测试夹具选中"}
                confirm={"event_type":"copy_confirmation","feedback_id":"confirm","batch_id":i["batch_id"],"candidate_id":i["candidate_id"],"revision_id":i["revision_id"],"format":mode,"content_sha256":"0"*64,"user_quote":"测试夹具确认"}
                rows=[selection,confirm]
                confirm["content_sha256"]=dc.content_hash(c)
                ledger.write_text("\n".join(json.dumps(row) for row in rows)+"\n",encoding="utf-8")
                source.write_text(json.dumps(e),encoding="utf-8")
                allowed=subprocess.run(cmd,capture_output=True,text=True,encoding="utf-8")
                self.assertEqual(0,allowed.returncode,allowed.stdout+allowed.stderr)
                words=list(root.glob("*.docx"));self.assertEqual(1,len(words));Document(words[0])
                before=registry.read_bytes();word_hash=dc.file_hash(words[0])
                again=subprocess.run(cmd,capture_output=True,text=True,encoding="utf-8")
                self.assertEqual(2,again.returncode)
                self.assertEqual(before,registry.read_bytes());self.assertEqual(word_hash,dc.file_hash(words[0]))

if __name__=="__main__":unittest.main()
